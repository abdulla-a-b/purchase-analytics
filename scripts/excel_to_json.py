#!/usr/bin/env python3
"""
Panda Purchase Analytics — Excel -> JSON data builder
=====================================================

Reads the input cells of the purchase workbook (Purchase_Register, Settings,
Budget_Control, Lists) and writes ``data/purchase_data.json``, the file the
web dashboard loads.  Only *inputs* are exported: every metric is recalculated
in the browser by ``assets/js/engine.js`` with the same rules as the Excel
formulas, so the website never depends on Excel having been recalculated.

Usage
-----
    python scripts/excel_to_json.py                       # default paths
    python scripts/excel_to_json.py --input my.xlsx --output data/purchase_data.json
    python scripts/excel_to_json.py --snapshot tests/excel_snapshot.json
        # also saves Excel's own calculated results, used by tests/engine.test.js
        # to prove the web engine and the workbook agree line by line.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

# Excel header text (row 5 of Purchase_Register) -> JSON key.  Matching is by
# header text, so columns can be moved without breaking the export.
INPUT_HEADERS = {
    "PR No": "pr_no", "PR Received Date": "pr_date", "Department": "dept", "Requested By": "req_by",
    "Item / Material Purchased": "item", "Category": "category", "Spec / Size / Colour": "spec",
    "Spend Type (OPEX / CAPEX)": "spend", "OPEX / CAPEX Approval Ref": "appr_ref", "Approval Date": "appr_date",
    "Approval Status": "appr_status", "Approved By": "appr_by", "Supplier": "supplier", "Local / Import": "origin",
    "PO No": "po_no", "PO Date": "po_date", "Order Quantity": "qty", "UOM": "uom",
    "Unit Price (PO Currency)": "unit_price", "Currency": "ccy", "Budget / Standard Unit Price (BDT)": "budget_price",
    "VAT / AIT / Duty %": "tax_pct", "Payment Terms": "pay_terms", "Initial / Advance Paid (BDT)": "adv_paid",
    "Interim Payment (BDT)": "interim_paid", "Closing Payment (BDT)": "closing_paid", "Last Payment Date": "last_pay_date",
    "Next Payment Due Date": "due_date", "Supplier Committed Delivery Date": "commit_date",
    "Starting to In-house Date": "inhouse_start", "Qty Received In-house": "qty_recv",
    "Final In-house Closing Date": "close_date", "QC / Material Approval": "qc_status",
    "Approved Material Photo — File / Link": "photo_link", "Remarks": "remarks",
}
DATE_KEYS = {"pr_date", "appr_date", "po_date", "last_pay_date", "due_date", "commit_date", "inhouse_start", "close_date"}
NUM_KEYS = {"qty", "unit_price", "budget_price", "tax_pct", "adv_paid", "interim_paid", "closing_paid", "qty_recv"}

# Calculated columns captured for the test snapshot (Excel's cached results).
SNAPSHOT_HEADERS = {
    "FX Rate to BDT": "fx", "PO Value (BDT)": "po_value", "Price Variance vs Budget (BDT)": "pv_bdt",
    "Price Variance %": "pv_pct", "Total Price incl. Tax (BDT)": "total_cost", "Total Paid (BDT)": "total_paid",
    "% of Payment": "pct_paid", "Outstanding Payable (BDT)": "outstanding", "Paid Ahead of Receipt (BDT)": "paid_ahead",
    "Finance Status": "fin_status", "Days Payment Overdue": "days_overdue", "Payable Aging Bucket": "aging",
    "Received %": "recv_pct", "Pending Qty to In-house": "pend_qty", "Pending In-house Value (BDT)": "pend_val",
    "Pending Material Status": "pend_status", "Days PR → Approval": "d_pr_appr", "Days Approval → PO": "d_appr_po",
    "Lead Time PO → In-house Close (days)": "d_po_close", "Total Cycle PR → Close (days)": "d_cycle",
    "Delivery Delay (days)": "delay", "OTIF": "otif", "Current Stage": "stage", "Unit Price (BDT)": "unit_bdt",
    "Previous PO Unit Price — Same Item (BDT)": "prev_price", "Price Change vs Previous PO %": "price_chg",
    "Data Check (auto)": "data_check", "Action Required (auto)": "action", "Action Priority": "priority",
}

SETTING_LABELS = {  # label text starts with ... -> key
    "Report Date": "report_date", "Company": "company", "OTIF quantity tolerance": "otif_tol",
    "Price variance alert threshold": "pv_alert", "Budget 'Watch' threshold": "watch", "Financial year start": "fy_start",
    "Approval SLA": "sla_appr", "PO issue SLA": "sla_po", "ABC class A cut-off": "abc_a", "ABC class B cut-off": "abc_b",
    "Quotations required": "min_quotes", "Quotation threshold": "quote_threshold", "Invoice price tolerance": "inv_tol",
}

# Log sheets (Deliveries, Payments, Quotations, Invoices): header text -> (key, type).
# Must match LOG_KINDS in assets/js/entry-rules.js — tests/logs.test.js checks it.
LOG_COMMON = {"ID": ("id", "text"), "PR No": ("pr_no", "text"), "Item / Material Purchased": ("item", "text"), "PO No": ("po_no", "text"),
              "Entered By": ("by", "text"), "Entered At": ("at", "datetime")}
LOG_SHEETS = {
    "deliveries": ("Deliveries", {"Receipt Date": ("date", "date"), "Qty Received": ("qty", "number"), "Qty Accepted at QC": ("qty_ok", "number"),
                                  "Delivery Challan No": ("challan", "text"), "Remarks": ("remarks", "text")}),
    "payments": ("Payments", {"Entry": ("entry", "text"), "Payment Type": ("ptype", "text"), "Date (paid or due)": ("date", "date"),
                              "Amount (BDT)": ("amount", "number"), "Invoice No": ("invoice_no", "text"), "Settles Planned ID": ("settles", "text"),
                              "Bank / Cheque Ref": ("ref", "text"), "Remarks": ("remarks", "text")}),
    "quotations": ("Quotations", {"Supplier": ("supplier", "text"), "Quote Date": ("date", "date"), "Unit Price": ("unit_price", "number"),
                                  "Currency": ("ccy", "text"), "Quoted Lead Time (days)": ("lead_days", "number"), "Payment Terms": ("pay_terms", "text"),
                                  "Remarks": ("remarks", "text")}),
    "invoices": ("Invoices", {"Supplier Invoice No": ("invoice_no", "text"), "Invoice Date": ("date", "date"), "Qty Invoiced": ("qty", "number"),
                              "Unit Price (PO Currency)": ("unit_price", "number"), "Invoice Amount incl. Tax (BDT)": ("amount", "number"),
                              "Remarks": ("remarks", "text")}),
}


def read_log(ws, spec, warnings):
    head = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    cols = {}
    for i, h in enumerate(head):
        if h in LOG_COMMON:
            cols[i] = LOG_COMMON[h]
        elif h in spec:
            cols[i] = spec[h]
    out = []
    for r in range(2, ws.max_row + 1):
        if all(ws.cell(r, i + 1).value in (None, "") for i in cols):
            continue
        e = {}
        for i, (k, typ) in cols.items():
            v = ws.cell(r, i + 1).value
            if v is None or (isinstance(v, str) and not v.strip()):
                e[k] = None
            elif typ == "date":
                e[k] = iso(v)
                if e[k] is None:
                    warnings.append(f"{ws.title} row {r} · {k}: '{v}' is not a date — ignored")
            elif typ == "datetime":
                e[k] = v.isoformat(timespec="minutes") if isinstance(v, dt.datetime) else str(v)
            elif typ == "number":
                try:
                    e[k] = float(v) if not isinstance(v, int) else v
                    if isinstance(e[k], float) and e[k].is_integer():
                        e[k] = int(e[k])
                except (TypeError, ValueError):
                    warnings.append(f"{ws.title} row {r} · {k}: '{v}' is not a number — ignored")
                    e[k] = None
            else:
                e[k] = str(v).strip()
        out.append(e)
    return out


def iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def clean(key, v, warnings, where):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if key in DATE_KEYS:
        d = iso(v)
        if d is None:
            warnings.append(f"{where}: '{v}' is not a date — ignored")
        return d
    if key in NUM_KEYS:
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            if v < 0:
                warnings.append(f"{where}: negative value {v}")
            return float(v) if isinstance(v, float) else int(v)
        try:
            return float(str(v).replace(",", ""))
        except ValueError:
            warnings.append(f"{where}: '{v}' is not a number — ignored")
            return None
    return str(v).strip()


def find_header_row(ws, must="PR No", max_row=15):
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            if ws.cell(r, c).value == must:
                return r
    raise SystemExit(f"Could not find the '{must}' header in sheet {ws.title}")


def read_settings(ws):
    out, fx, in_fx = {}, {}, False
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a is None:
            in_fx = False
            continue
        a = str(a)
        if a == "Currency" and str(b).startswith("Rate"):
            in_fx = True
            continue
        if in_fx and isinstance(b, (int, float)):
            fx[a.strip()] = float(b)
            continue
        for label, key in SETTING_LABELS.items():
            if a.startswith(label):
                out[key] = iso(b) if key in ("report_date", "fy_start") else b
    out["fx"] = fx
    return out


def read_budgets(ws):
    hr = find_header_row(ws, "Department")
    head = {ws.cell(hr, c).value: c for c in range(1, ws.max_column + 1)}
    out = []
    for r in range(hr + 1, ws.max_row + 1):
        d = ws.cell(r, head["Department"]).value
        if not d or d == "TOTAL":
            continue
        out.append({"dept": str(d).strip(),
                    "opex": ws.cell(r, head["OPEX Budget"]).value or 0,
                    "capex": ws.cell(r, head["CAPEX Budget"]).value or 0})
    return out


def read_lists(ws):
    keymap = {"Department": "dept", "Category": "category", "Spend Type": "spend", "Approval Status": "appr_status",
              "Origin": "origin", "UOM": "uom", "Currency": "ccy", "Payment Terms": "pay_terms",
              "QC / Material Approval": "qc_status", "Supplier": "supplier"}
    out = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h in keymap:
            vals = [ws.cell(r, c).value for r in range(2, 42)]
            out[keymap[h]] = [str(v).strip() for v in vals if v not in (None, "")]
    return out


def build(path: Path, snapshot: Path | None = None) -> dict:
    wb = load_workbook(path, data_only=True)
    reg = wb["Purchase_Register"]
    hr = find_header_row(reg)
    cols = {reg.cell(hr, c).value: c for c in range(1, reg.max_column + 1) if reg.cell(hr, c).value}
    missing = [h for h in INPUT_HEADERS if h not in cols]
    if missing:
        raise SystemExit("Purchase_Register is missing columns: " + ", ".join(missing))

    warnings, lines, snap = [], [], []
    for r in range(hr + 1, reg.max_row + 1):
        if reg.cell(r, cols["PR No"]).value in (None, ""):
            continue
        line = {"_row": r}
        for h, k in INPUT_HEADERS.items():
            line[k] = clean(k, reg.cell(r, cols[h]).value, warnings, f"Row {r} · {h}")
        lines.append(line)
        if snapshot is not None:
            s = {"_row": r, "pr_no": line["pr_no"]}
            for h, k in SNAPSHOT_HEADERS.items():
                if h in cols:
                    v = reg.cell(r, cols[h]).value
                    s[k] = None if v == "" else v
            snap.append(s)

    settings = read_settings(wb["Settings"])
    data = {
        "meta": {
            "schema": 1,
            "company": settings.pop("company", "Panda Shoes Industry Ltd"),
            "source": path.name,
            "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
            "excel_report_date": settings.pop("report_date", None),
            "rows": len(lines),
            "warnings": warnings,
        },
        "settings": settings,
        "budgets": read_budgets(wb["Budget_Control"]),
        "lists": read_lists(wb["Lists"]),
        "lines": lines,
    }
    for kind, (sheet, spec) in LOG_SHEETS.items():
        if sheet in wb.sheetnames:
            data[kind] = read_log(wb[sheet], spec, warnings)
    if snapshot is not None:
        snapshot.parent.mkdir(parents=True, exist_ok=True)
        sc = wb["Supplier_Scorecard"]
        suppliers = [{"supplier": sc.cell(r, 1).value, "spend": sc.cell(r, 3).value, "otif": sc.cell(r, 8).value,
                      "infull": sc.cell(r, 9).value, "qc": sc.cell(r, 15).value, "score": sc.cell(r, 17).value,
                      "rating": sc.cell(r, 18).value}
                     for r in range(5, 35) if sc.cell(r, 1).value]
        sp = wb["Spend_Analysis"]
        cats = [{"category": sp.cell(r, 2).value, "spend": sp.cell(r, 4).value, "cls": sp.cell(r, 8).value}
                for r in range(6, 26) if sp.cell(r, 2).value]
        bc = wb["Budget_Control"]
        buds = [{"dept": bc.cell(r, 1).value, "committed": bc.cell(r, 9).value, "status": bc.cell(r, 13).value}
                for r in range(5, 25) if bc.cell(r, 1).value]
        snapshot.write_text(json.dumps({"report_date": data["meta"]["excel_report_date"], "lines": snap,
                                        "suppliers": suppliers, "categories": cats, "budgets": buds},
                                       ensure_ascii=False, indent=1, default=str))
    return data


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=str(ROOT / "data" / "Panda_Purchase_Tracking_Dashboard.xlsx"))
    ap.add_argument("--output", default=str(ROOT / "data" / "purchase_data.json"))
    ap.add_argument("--snapshot", default=None, help="also write Excel's calculated values for the test-suite")
    a = ap.parse_args(argv)
    data = build(Path(a.input), Path(a.snapshot) if a.snapshot else None)
    out = Path(a.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=1))
    print(f"✔ {data['meta']['rows']} purchase lines -> {out}")
    for w in data["meta"]["warnings"]:
        print("  ⚠", w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
